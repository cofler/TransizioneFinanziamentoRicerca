"""Conversione di MODELLO_spiegazione.md in .docx.

Non e' un convertitore Markdown generico: copre esattamente il sottoinsieme usato dal
documento - titoli, paragrafi, liste, tabelle con pipe, citazioni (>), blocchi di codice,
e l'inline grassetto/corsivo/codice/link. Tutto il resto passa come testo.

Sta nel repo e non nello scratchpad perche' il .docx e' un deliverable: chi rigenera il
modello deve poter rigenerare anche il documento senza rimettere in piedi il convertitore.

Uso:  python md2docx.py [sorgente.md] [destinazione.docx]
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

QUI = Path(__file__).parent
SRC = QUI / "MODELLO_spiegazione.md"
DST = QUI / "MODELLO_spiegazione.docx"

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x55, 0x55, 0x55)
ACCENT = RGBColor(0x1F, 0x3B, 0x73)

# grassetto, corsivo, codice, link: l'ordine conta - i link per primi, perche' il
# testo dell'etichetta puo' contenere a sua volta grassetto
INLINE = re.compile(
    r"(?P<link>\[(?P<lab>[^\]]+)\]\((?P<url>[^)]+)\))"
    r"|(?P<bold>\*\*(?P<btxt>.+?)\*\*)"
    r"|(?P<code>`(?P<ctxt>[^`]+)`)"
    r"|(?P<ital>(?<![*\w])\*(?P<itxt>[^*]+?)\*(?![*\w]))"
)


def _sfondo(cella, esa: str) -> None:
    """Colore di sfondo di una cella: python-docx non lo espone, si scende all'XML."""
    tcPr = cella._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), esa)
    tcPr.append(shd)


def _barra(par, esa: str = "C8CCD4") -> None:
    """Barra verticale a sinistra del paragrafo, per rendere le citazioni."""
    pPr = par._p.get_or_add_pPr()
    bordi = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), esa)
    bordi.append(left)
    pPr.append(bordi)


def _scrivi_inline(par, testo: str, *, base_bold=False, corsivo=False,
                   colore=None) -> None:
    """Aggiunge `testo` al paragrafo interpretando la formattazione inline."""
    pos = 0
    for m in INLINE.finditer(testo):
        if m.start() > pos:
            r = par.add_run(testo[pos:m.start()])
            r.bold, r.italic = base_bold, corsivo
            if colore:
                r.font.color.rgb = colore
        if m.group("link"):
            r = par.add_run(m.group("lab"))
            r.font.color.rgb = ACCENT
            r.underline = True
            r.bold = base_bold
        elif m.group("bold"):
            r = par.add_run(m.group("btxt"))
            r.bold = True
            r.italic = corsivo
            if colore:
                r.font.color.rgb = colore
        elif m.group("code"):
            r = par.add_run(m.group("ctxt"))
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.bold = base_bold
        else:
            r = par.add_run(m.group("itxt"))
            r.italic = True
            r.bold = base_bold
            if colore:
                r.font.color.rgb = colore
        pos = m.end()
    if pos < len(testo):
        r = par.add_run(testo[pos:])
        r.bold, r.italic = base_bold, corsivo
        if colore:
            r.font.color.rgb = colore


def _riga_tabella(linea: str) -> list[str]:
    return [c.strip() for c in linea.strip().strip("|").split("|")]


def _e_separatore(linea: str) -> bool:
    return bool(re.fullmatch(r"\s*\|?[\s:|-]+\|?\s*", linea)) and "-" in linea


def _tabella(doc: Document, blocco: list[str]) -> None:
    righe = [_riga_tabella(l) for l in blocco if not _e_separatore(l)]
    if not righe:
        return
    ncol = max(len(r) for r in righe)
    t = doc.add_table(rows=len(righe), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, riga in enumerate(righe):
        for j in range(ncol):
            cella = t.cell(i, j)
            cella.text = ""
            par = cella.paragraphs[0]
            par.paragraph_format.space_before = Pt(2)
            par.paragraph_format.space_after = Pt(2)
            testo = riga[j] if j < len(riga) else ""
            _scrivi_inline(par, testo, base_bold=(i == 0))
            for r in par.runs:
                r.font.size = Pt(8.5)
            if i == 0:
                _sfondo(cella, "E8EAEE")
    doc.add_paragraph()


def _codice(doc: Document, blocco: list[str]) -> None:
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Cm(0.5)
    par.paragraph_format.space_before = Pt(4)
    par.paragraph_format.space_after = Pt(8)
    r = par.add_run("\n".join(blocco))
    r.font.name = "Consolas"
    r.font.size = Pt(9)


def converti(src: Path, dst: Path) -> None:
    doc = Document()
    normale = doc.styles["Normal"]
    normale.font.name = "Calibri"
    normale.font.size = Pt(10.5)
    normale.font.color.rgb = INK
    normale.paragraph_format.space_after = Pt(6)
    normale.paragraph_format.line_spacing = 1.15

    righe = src.read_text(encoding="utf-8").splitlines()
    i, buf_tab, buf_cod, in_codice = 0, [], [], False

    def chiudi_tabella():
        nonlocal buf_tab
        if buf_tab:
            _tabella(doc, buf_tab)
            buf_tab = []

    while i < len(righe):
        linea = righe[i]

        if linea.strip().startswith("```"):
            if in_codice:
                _codice(doc, buf_cod)
                buf_cod, in_codice = [], False
            else:
                chiudi_tabella()
                in_codice = True
            i += 1
            continue
        if in_codice:
            buf_cod.append(linea)
            i += 1
            continue

        # tabella: una riga con pipe che non sia una citazione
        if linea.strip().startswith("|") and linea.strip().endswith("|"):
            buf_tab.append(linea)
            i += 1
            continue
        chiudi_tabella()

        nudo = linea.strip()

        if not nudo:
            i += 1
            continue

        if re.fullmatch(r"-{3,}", nudo):        # regola orizzontale -> spazio
            doc.add_paragraph()
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", nudo)
        if m:
            liv = len(m.group(1))
            par = doc.add_heading("", level=min(liv, 4))
            _scrivi_inline(par, m.group(2))
            for r in par.runs:
                r.font.color.rgb = ACCENT if liv <= 2 else INK
            i += 1
            continue

        if nudo.startswith(">"):
            # una citazione puo' contenere una tabella: si raccoglie il blocco intero
            blocco = []
            while i < len(righe) and righe[i].strip().startswith(">"):
                blocco.append(re.sub(r"^\s*>\s?", "", righe[i]))
                i += 1
            sotto_tab = [l for l in blocco if l.strip().startswith("|")]
            testo = [l for l in blocco if not l.strip().startswith("|")]
            unito = " ".join(l.strip() for l in testo if l.strip())
            if unito:
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Cm(0.4)
                par.paragraph_format.space_before = Pt(6)
                _barra(par)
                _scrivi_inline(par, unito, colore=MUTED)
                for r in par.runs:
                    r.font.size = Pt(10)
            if sotto_tab:
                _tabella(doc, sotto_tab)
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", nudo)
        if m:
            par = doc.add_paragraph(style="List Number")
            _scrivi_inline(par, m.group(2))
            i += 1
            continue

        if re.match(r"^[-*]\s+", nudo):
            par = doc.add_paragraph(style="List Bullet")
            _scrivi_inline(par, re.sub(r"^[-*]\s+", "", nudo))
            i += 1
            continue

        # paragrafo: si accorpano le righe successive fino alla riga vuota, perche'
        # nel sorgente il testo e' mandato a capo a 88 colonne
        blocco = [nudo]
        i += 1
        while i < len(righe):
            succ = righe[i].strip()
            if (not succ or succ.startswith(("#", ">", "|", "```", "- ", "* "))
                    or re.match(r"^\d+\.\s", succ) or re.fullmatch(r"-{3,}", succ)):
                break
            blocco.append(succ)
            i += 1
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _scrivi_inline(par, " ".join(blocco))

    chiudi_tabella()
    doc.save(dst)
    print(f"scritto {dst.name} ({dst.stat().st_size/1024:.0f} KB)")


if __name__ == "__main__":
    s = Path(sys.argv[1]) if len(sys.argv) > 1 else SRC
    d = Path(sys.argv[2]) if len(sys.argv) > 2 else DST
    converti(s, d)
